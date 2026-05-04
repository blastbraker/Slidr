"""Template builder v3 - 10 slide layouts with image embedding"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from typing import List, Dict, Any, Optional
import os
import tempfile
import requests
from io import BytesIO


SLIDE_TYPES = {
    "title": "Title Slide - Big title with subtitle",
    "content": "Content - Title and bullet points",
    "bullets_image": "Bullets + Image - Title, bullets on left, image on right",
    "two_column": "Two Column - Two content areas side by side",
    "divider": "Divider - Big centered title for section transitions",
    "quote": "Quote - Big quote with author citation",
    "statistic": "Statistic - Large number with label",
    "comparison": "Comparison - Two-column pros/cons table",
    "timeline": "Timeline - Chronological events",
    "full_image": "Full Image - Hero image with caption overlay",
}


class TemplateBuilder:
    """Professional templates with 10 layouts"""

    THEMES = {
        "minimal": {"primary": "1E3A5F", "accent": "E74C3C"},
        "modern": {"primary": "1A1A2E", "accent": "E94560"},
        "corporate": {"primary": "2C3E50", "accent": "27AE60"},
    }

    def __init__(self, theme: str = "minimal"):
        self.theme = theme
        self.colors = self.THEMES.get(theme, self.THEMES["minimal"])

    def _rgb(self, h: str) -> RGBColor:
        h = h.lstrip("#")
        return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def _download_image(self, url: str) -> Optional[bytes]:
        """Download image from URL"""
        if not url or not url.startswith("http"):
            return None
        try:
            resp = requests.get(url, timeout=10)
            if resp.status_code == 200:
                return resp.content
        except:
            pass
        return None

    def _add_image(self, slide, url: str, left: float, top: float, width: float, height: float) -> bool:
        """Add image to slide, returns True if successful"""
        img_data = self._download_image(url)
        if img_data:
            try:
                with tempfile.NamedTemporaryFile(suffix=".jpg", delete=False) as f:
                    f.write(img_data)
                    f.flush()
                    img_path = f.name
                try:
                    slide.shapes.add_picture(img_path, Inches(left), Inches(top), Inches(width), Inches(height))
                    os.unlink(img_path)
                    return True
                except:
                    if os.path.exists(img_path):
                        os.unlink(img_path)
            except:
                pass
        return False

    def create_presentation(self, slides: List[Dict[str, Any]], title: str) -> Presentation:
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        for i, slide_data in enumerate(slides):
            slide_type = slide_data.get("type", "content")
            content_slide = prs.slides.add_slide(prs.slide_layouts[6])

            if slide_type == "title":
                self._layout_title(content_slide, slide_data)
            elif slide_type == "divider":
                self._layout_divider(content_slide, slide_data)
            elif slide_type == "bullets_image":
                self._layout_bullets_image(content_slide, slide_data)
            elif slide_type == "two_column":
                self._layout_two_column(content_slide, slide_data)
            elif slide_type == "quote":
                self._layout_quote(content_slide, slide_data)
            elif slide_type == "statistic":
                self._layout_statistic(content_slide, slide_data)
            elif slide_type == "comparison":
                self._layout_comparison(content_slide, slide_data)
            elif slide_type == "timeline":
                self._layout_timeline(content_slide, slide_data)
            elif slide_type == "full_image":
                self._layout_full_image(content_slide, slide_data)
            else:
                self._layout_content(content_slide, slide_data)

        return prs

    def _add_background(self, shapes, primary_color: str = None):
        c = self.colors
        bg = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = self._rgb(primary_color or c["primary"])
        bg.line.fill.background()
        return bg

    def _add_header(self, shapes, title: str, font_size: int = 28):
        c = self.colors
        hdr = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.9))
        hdr.fill.solid()
        hdr.fill.fore_color.rgb = self._rgb(c["accent"])
        hdr.line.fill.background()

        tb = shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(font_size)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        return tb

    def _layout_title(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)

        bar = s.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.1), Inches(7.5))
        bar.fill.solid()
        bar.fill.fore_color.rgb = self._rgb(c["accent"])
        bar.line.fill.background()

        tb = s.add_textbox(Inches(1.5), Inches(2.2), Inches(10), Inches(2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data.get("title", "Presentation")
        p.font.size = Pt(52)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)

        sub = data.get("subtitle", "")
        if sub:
            sb = s.add_textbox(Inches(1.5), Inches(4.2), Inches(8), Inches(0.8))
            tf = sb.text_frame
            p = tf.paragraphs[0]
            p.text = sub
            p.font.size = Pt(24)
            p.font.color.rgb = self._rgb(c["accent"])

        image_url = data.get("image_url", "")
        if image_url:
            self._add_image(slide, image_url, 9, 1.5, 3.5, 4)

    def _layout_content(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)
        self._add_header(s, data.get("title", "Slide"))

        sub = data.get("subtitle", "")
        y = 1.5 if sub else 1.2
        if sub:
            sb = s.add_textbox(Inches(0.5), Inches(1.0), Inches(12), Inches(0.5))
            tf = sb.text_frame
            p = tf.paragraphs[0]
            p.text = sub
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(200, 200, 200)

        bb = s.add_textbox(Inches(0.5), Inches(y), Inches(12), Inches(5))
        tf = bb.text_frame
        tf.word_wrap = True

        for i, b in enumerate(data.get("bullets", [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(20)
            p.level = 0
            p.space_before = Pt(12)
            p.font.color.rgb = RGBColor(230, 230, 230)

    def _layout_bullets_image(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)
        self._add_header(s, data.get("title", "Slide"))

        bb = s.add_textbox(Inches(0.5), Inches(1.1), Inches(5.5), Inches(5))
        tf = bb.text_frame
        tf.word_wrap = True

        for i, b in enumerate(data.get("bullets", [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(18)
            p.level = 0
            p.space_before = Pt(10)
            p.font.color.rgb = RGBColor(230, 230, 230)

        image_url = data.get("image_url", "")
        kw = data.get("image_keywords", "")

        if image_url and image_url.startswith("http"):
            if not self._add_image(slide, image_url, 6.2, 1.1, 6.5, 5.2):
                img_text = f"[Image: {kw}]" if kw else "[No image]"
        elif kw:
            img_text = f"[Image: {kw}]"
        else:
            img_text = "[Image placeholder]"

        if image_url and image_url.startswith("http"):
            self._add_image(slide, image_url, 6.2, 1.1, 6.5, 5.2)
        else:
            img_box = s.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.2), Inches(1.1), Inches(6.5), Inches(5.2))
            img_box.fill.solid()
            img_box.fill.fore_color.rgb = RGBColor(60, 60, 60)
            img_box.line.color.rgb = self._rgb(c["accent"])

            tb = s.add_textbox(Inches(6.2), Inches(3.5), Inches(6.5), Inches(1))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = f"[Image: {kw}]" if kw else "[Image placeholder]"
            p.font.size = Pt(16)
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = RGBColor(150, 150, 150)

    def _layout_two_column(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)
        self._add_header(s, data.get("title", "Slide"))

        left = s.add_textbox(Inches(0.5), Inches(1.1), Inches(5.8), Inches(5))
        tf = left.text_frame
        tf.word_wrap = True

        bullets = data.get("bullets", [])
        left_bullets = bullets[:3] if bullets else ["Point 1", "Point 2"]

        for i, b in enumerate(left_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(18)
            p.level = 0
            p.space_before = Pt(10)
            p.font.color.rgb = RGBColor(230, 230, 230)

        div = s.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.5), Inches(1.1), Inches(0.03), Inches(5))
        div.fill.solid()
        div.fill.fore_color.rgb = self._rgb(c["accent"])
        div.line.fill.background()

        right = s.add_textbox(Inches(7), Inches(1.1), Inches(5.8), Inches(5))
        tf = right.text_frame
        tf.word_wrap = True

        right_bullets = bullets[3:] if len(bullets) > 3 else ["Point 4", "Point 5"]

        for i, b in enumerate(right_bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(18)
            p.level = 0
            p.space_before = Pt(10)
            p.font.color.rgb = RGBColor(230, 230, 230)

    def _layout_divider(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s, c["accent"])

        tb = s.add_textbox(Inches(1), Inches(2.8), Inches(11.333), Inches(2.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data.get("title", "Section")
        p.font.size = Pt(56)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        p.font.color.rgb = RGBColor(255, 255, 255)

        sub = data.get("subtitle", "")
        if sub:
            sb = s.add_textbox(Inches(1), Inches(5), Inches(11.333), Inches(1))
            tf = sb.text_frame
            p = tf.paragraphs[0]
            p.text = sub
            p.font.size = Pt(24)
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = RGBColor(255, 255, 255)

    def _layout_quote(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)

        quote = data.get("quote", "")
        if not quote and data.get("bullets"):
            quote = data.get("bullets", [""])[0]

        tb = s.add_textbox(Inches(1.5), Inches(1.5), Inches(10), Inches(3.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f'"{quote}"' if quote else '"Quote"'
        p.font.size = Pt(36)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        p.font.color.rgb = RGBColor(255, 255, 255)

        author = data.get("author", "")
        if author:
            ab = s.add_textbox(Inches(1.5), Inches(5.2), Inches(10), Inches(0.8))
            tf = ab.text_frame
            p = tf.paragraphs[0]
            p.text = f"— {author}"
            p.font.size = Pt(20)
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = self._rgb(c["accent"])
        else:
            ab = s.add_textbox(Inches(1.5), Inches(5.2), Inches(10), Inches(0.8))
            tf = ab.text_frame
            p = tf.paragraphs[0]
            p.text = data.get("title", "")
            p.font.size = Pt(18)
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = self._rgb(c["accent"])

    def _layout_statistic(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)
        self._add_header(s, data.get("title", "Statistics"))

        big_num = data.get("big_number", "")
        stat_label = data.get("stat_label", "")

        if not big_num and data.get("bullets"):
            parts = data.get("bullets", [""])[0].split()
            if parts:
                big_num = parts[0]
                stat_label = " ".join(parts[1:]) if len(parts) > 1 else "Growth"

        tb = s.add_textbox(Inches(1), Inches(2), Inches(11.333), Inches(2.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = big_num if big_num else "85%"
        p.font.size = Pt(72)
        p.font.bold = True
        p.alignment = PP_ALIGN.CENTER
        p.font.color.rgb = RGBColor(255, 255, 255)

        if stat_label:
            sb = s.add_textbox(Inches(1), Inches(4.5), Inches(11.333), Inches(1))
            tf = sb.text_frame
            p = tf.paragraphs[0]
            p.text = stat_label if stat_label else "Increase"
            p.font.size = Pt(28)
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = self._rgb(c["accent"])

        sub = data.get("subtitle", "")
        if sub:
            sb = s.add_textbox(Inches(1), Inches(5.5), Inches(11.333), Inches(1))
            tf = sb.text_frame
            p = tf.paragraphs[0]
            p.text = sub
            p.font.size = Pt(16)
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = RGBColor(200, 200, 200)

    def _layout_comparison(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)
        self._add_header(s, data.get("title", "Comparison"))

        left_title = data.get("left_title", "Pros")
        right_title = data.get("right_title", "Cons")
        left_items = data.get("left_items", data.get("bullets", [])[:3])
        right_items = data.get("right_items", data.get("bullets", [])[3:])

        lt = s.add_textbox(Inches(0.5), Inches(1.1), Inches(6), Inches(0.6))
        tf = lt.text_frame
        p = tf.paragraphs[0]
        p.text = left_title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = self._rgb(c["accent"])

        rt = s.add_textbox(Inches(7), Inches(1.1), Inches(6), Inches(0.6))
        tf = rt.text_frame
        p = tf.paragraphs[0]
        p.text = right_title
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 100, 100)

        left = s.add_textbox(Inches(0.5), Inches(1.8), Inches(6), Inches(5))
        tf = left.text_frame
        tf.word_wrap = True

        for i, b in enumerate(left_items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(18)
            p.space_before = Pt(10)
            p.font.color.rgb = RGBColor(230, 230, 230)

        right = s.add_textbox(Inches(7), Inches(1.8), Inches(6), Inches(5))
        tf = right.text_frame
        tf.word_wrap = True

        for i, b in enumerate(right_items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = b
            p.font.size = Pt(18)
            p.space_before = Pt(10)
            p.font.color.rgb = RGBColor(230, 230, 230)

    def _layout_timeline(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes
        self._add_background(s)
        self._add_header(s, data.get("title", "Timeline"))

        events = data.get("events", [])
        if not events and data.get("bullets"):
            events = [{"title": b} for b in data.get("bullets", [])]

        y_pos = 1.8
        for i, event in enumerate(events[:5]):
            date = event.get("date", "")
            title = event.get("title", event.get("text", ""))

            marker = s.add_shape(MSO_SHAPE.OVAL, Inches(1), Inches(y_pos), Inches(0.3), Inches(0.3))
            marker.fill.solid()
            marker.fill.fore_color.rgb = self._rgb(c["accent"])
            marker.line.color.rgb = self._rgb(c["accent"])

            if date:
                tb = s.add_textbox(Inches(1.5), Inches(y_pos), Inches(2), Inches(0.4))
                tf = tb.text_frame
                p = tf.paragraphs[0]
                p.text = date
                p.font.size = Pt(14)
                p.font.bold = True
                p.font.color.rgb = self._rgb(c["accent"])

            tb = s.add_textbox(Inches(3.5), Inches(y_pos), Inches(9), Inches(0.4))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(230, 230, 230)

            if i < len(events) - 1:
                line = s.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.1), Inches(y_pos + 0.35), Inches(0.1), Inches(0.4))
                line.fill.solid()
                line.fill.fore_color.rgb = self._rgb(c["accent"])
                line.line.fill.background()

            y_pos += 1.1

    def _layout_full_image(self, slide, data: Dict[str, Any]):
        c = self.colors
        s = slide.shapes

        image_url = data.get("image_url", "")
        kw = data.get("image_keywords", "")

        if image_url and image_url.startswith("http"):
            if not self._add_image(slide, image_url, 0, 0, 13.333, 7.5):
                bg = s.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
                bg.fill.solid()
                bg.fill.fore_color.rgb = RGBColor(40, 40, 40)
                bg.line.fill.background()
        else:
            bg = s.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
            bg.fill.solid()
            bg.fill.fore_color.rgb = RGBColor(40, 40, 40)
            bg.line.fill.background()

            tb = s.add_textbox(Inches(3), Inches(3.5), Inches(7.333), Inches(1))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = f"[Image: {kw}]" if kw else "[Image placeholder]"
            p.font.size = Pt(24)
            p.alignment = PP_ALIGN.CENTER
            p.font.color.rgb = RGBColor(150, 150, 150)

        overlay = data.get("overlay_title", "")
        if overlay:
            tb = s.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.7))
            tf = tb.text_frame
            p = tf.paragraphs[0]
            p.text = overlay
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)

        caption = data.get("caption", "")
        if caption:
            cb = s.add_textbox(Inches(0.5), Inches(6.8), Inches(12), Inches(0.5))
            tf = cb.text_frame
            p = tf.paragraphs[0]
            p.text = caption
            p.font.size = Pt(14)
            p.font.color.rgb = RGBColor(200, 200, 200)


def create_from_template(slides: List[Dict[str, Any]], title: str, theme: str = "minimal", template_path: Optional[str] = None) -> Presentation:
    if template_path and os.path.exists(template_path):
        return Presentation(template_path)
    return TemplateBuilder(theme).create_presentation(slides, title)


if __name__ == "__main__":
    test = [
        {"type": "title", "title": "Introduction to AI", "subtitle": "Exploring the Future", "bullets": [], "image_keywords": "artificial intelligence"},
        {"type": "content", "title": "What is AI?", "subtitle": "Understanding artificial intelligence", "bullets": ["Machine learning", "Neural networks", "Deep learning"]},
        {"type": "statistic", "big_number": "85%", "stat_label": "Accuracy", "title": "Performance"},
        {"type": "comparison", "title": "Pros & Cons", "left_title": "Pros", "left_items": ["Fast", "Accurate"], "right_title": "Cons", "right_items": ["Cost", "Complexity"]},
        {"type": "timeline", "title": "History", "events": [{"date": "1950", "title": "AI Beginnings"}, {"date": "2020", "title": "Modern AI"}]},
        {"type": "quote", "quote": "AI will change everything", "author": "Expert"},
        {"type": "bullets_image", "title": "Applications", "bullets": ["Healthcare", "Finance"]},
        {"type": "two_column", "title": "Benefits", "bullets": ["Efficiency", "Accuracy", "Speed", "Cost"]},
        {"type": "divider", "title": "Conclusion"},
        {"type": "full_image", "caption": "AI Technology", "overlay_title": "Future of AI"},
    ]
    for t in ["minimal", "modern", "corporate"]:
        TemplateBuilder(t).create_presentation(test, f"Test {t}").save(f"test_{t}.pptx")
    print("Created test files")